"""Planilha com o preço de cada nível de cada encantamento: economia/saida/Precos_por_nivel.xlsx

O preço de cada nível é o que o plugin calcula a partir das faixas do encantamentos.yml
(custo no começo da faixa + aumento x níveis desde o começo). Uma aba por encantamento.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

ECON = Path(__file__).resolve().parents[1]
if str(ECON) not in sys.path:
    sys.path.insert(0, str(ECON))

from sim.modelo import Modelo, BOOST, AREA, FIXOS, CHAVE_PLUGIN   # noqa: E402
from gen.planilha import NOME_ENC                                 # noqa: E402
from gen.encantamentos import chance_area                         # noqa: E402

CAB = Font(bold=True, color='FFFFFF')
FUNDO = PatternFill('solid', fgColor='1F4E78')
ZEBRA = PatternFill('solid', fgColor='F2F2F2')
FINA = Side(style='thin', color='BFBFBF')
BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)


def br(v):
    return f'{v:,}'.replace(',', '.')


def aba(wb, titulo, linhas, cab, larguras):
    ws = wb.create_sheet(titulo[:31])
    for j, c in enumerate(cab, 1):
        cel = ws.cell(1, j, c)
        cel.font, cel.fill, cel.border = CAB, FUNDO, BORDA
        cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 32
    for i, linha in enumerate(linhas, 2):
        for j, v in enumerate(linha, 1):
            cel = ws.cell(i, j, v)
            cel.border = BORDA
            if isinstance(v, int):
                cel.number_format = '#,##0'
                cel.alignment = Alignment(horizontal='right')
            if i % 2:
                cel.fill = ZEBRA
    for j, w in enumerate(larguras, 1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = 'A2'
    return ws


def gerar():
    mo = Modelo()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Como ler'
    texto = [
        'Preço de cada nível dos encantamentos da picareta',
        '',
        'Uma aba por encantamento. Cada linha é um upgrade: "do nível 0 para o 1", "do 1 para o 2"...',
        'Money e Gemas = quanto custa aquele upgrade.',
        'Fortuna, Gemado e Chaveiro: a 2a coluna é o bônus (%) que o encantamento dá no nível novo.',
        'Encantamentos de área: a 2a coluna é a chance de ativar (%) a cada bloco quebrado, no nível novo.',
        'Total acumulado = quanto custa sair do nível 0 até aquele nível.',
        '',
        'Estes valores NÃO precisam ser digitados um por um. No encantamentos.yml ficam só as faixas',
        '(coluna "Faixa no arquivo"): o plugin faz a conta sozinho (custo do começo da faixa + aumento a cada nível).',
        'Esta planilha serve para vocês conferirem: o menu do jogo deve mostrar exatamente estes preços.',
        '',
        'Nos encantamentos de área, o upgrade do nível 0 para o 1 é mais caro: é o preço de liberar o encantamento.',
    ]
    for i, t in enumerate(texto, 1):
        ws.cell(i, 1, t).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    ws.column_dimensions['A'].width = 110

    pct_boost = {'fortuna': mo.P['picareta.fortuna_pct_por_nivel'], 'gemado': mo.P['picareta.gemado_pct_por_nivel'],
                 'chaveiro': mo.P['design.chaveiro_pct_por_nivel']}
    nome_efeito = {'fortuna': 'Bônus de money no nível novo (%)', 'gemado': 'Bônus de gemas no nível novo (%)',
                   'chaveiro': 'Bônus de chaves no nível novo (%)'}
    for e in BOOST + AREA:
        faixas = mo.faixas_preco(e)
        tab = mo.tabela_precos(e)
        if e in BOOST:
            por_nivel, titulo, casas = pct_boost[e], nome_efeito[e], 1
        else:
            por_nivel, titulo, casas = chance_area(mo, e), 'Chance de ativação no nível novo (% por bloco)', 8
        linhas, acm, acg = [], 0, 0
        for a, b, *_ in faixas:
            for n in range(a, b + 1):
                m, g = tab[n]
                acm += m
                acg += g
                efeito = round(por_nivel * (n + 1), casas)
                linhas.append([f'{n} → {n + 1}', efeito, int(m), int(g), f"'{a}-{b}'", int(acm), int(acg)])
        ws = aba(wb, NOME_ENC[e], linhas,
                 ['Upgrade (nível)', titulo, 'Money', 'Gemas', 'Faixa no arquivo', 'Total acumulado money',
                  'Total acumulado gemas'],
                 [16, 24, 18, 16, 17, 24, 24])
        for linha in ws.iter_rows(min_row=2, min_col=2, max_col=2):
            linha[0].number_format = '0.0' if e in BOOST else '0.00000000'
            linha[0].alignment = Alignment(horizontal='right')
    linhas = []
    for e in FIXOS:
        for L in (1, 2):
            m, g = mo.preco(e, L)
            linhas.append([NOME_ENC[e], f'{L - 1} → {L}', int(m), int(g)])
    aba(wb, 'Speed e Haste', linhas, ['Encantamento', 'Upgrade (nível)', 'Money', 'Gemas'], [16, 16, 18, 16])
    destino = ECON / 'saida' / 'Precos_por_nivel.xlsx'
    wb.save(destino)
    return destino


if __name__ == '__main__':
    print(gerar())
